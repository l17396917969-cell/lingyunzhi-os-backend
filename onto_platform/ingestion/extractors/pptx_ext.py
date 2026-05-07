import asyncio

from pptx import Presentation

from onto_platform.ingestion.extractors.types import (
    ExtractedDocument,
    ExtractionFailure,
    Section,
)


async def extract_pptx(path: str, filename: str) -> ExtractedDocument:
    def _do() -> ExtractedDocument:
        try:
            pres = Presentation(path)
        except Exception as e:
            raise ExtractionFailure(f"python-pptx failed on {filename}: {e}") from e
        sections: list[Section] = []
        for i, slide in enumerate(pres.slides):
            heading = ""
            body_lines: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text:
                    if slide.shapes.title is not None and shape == slide.shapes.title:
                        heading = shape.text_frame.text
                    else:
                        body_lines.append(shape.text_frame.text)
            sections.append(
                Section(
                    heading=heading or f"slide {i + 1}",
                    text="\n".join(body_lines),
                )
            )
        return ExtractedDocument(kind="pptx", source_filename=filename, sections=sections)

    return await asyncio.to_thread(_do)
