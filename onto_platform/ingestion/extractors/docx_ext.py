import asyncio

from docx import Document

from onto_platform.ingestion.extractors.types import (
    ExtractedDocument,
    ExtractionFailure,
    Section,
)


async def extract_docx(path: str, filename: str) -> ExtractedDocument:
    def _do() -> ExtractedDocument:
        try:
            d = Document(path)
        except Exception as e:
            raise ExtractionFailure(f"python-docx failed on {filename}: {e}") from e
        sections: list[Section] = []
        current_heading = ""
        current_text: list[str] = []

        def flush() -> None:
            if current_heading or current_text:
                sections.append(
                    Section(heading=current_heading, text="\n".join(current_text))
                )

        for para in d.paragraphs:
            style = ((para.style.name if para.style is not None else None) or "").lower()
            if style.startswith("heading"):
                flush()
                current_heading = para.text
                current_text.clear()
            else:
                current_text.append(para.text)
        flush()
        return ExtractedDocument(kind="docx", source_filename=filename, sections=sections)

    return await asyncio.to_thread(_do)
