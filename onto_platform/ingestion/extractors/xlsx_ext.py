import asyncio
from typing import Any

from openpyxl import load_workbook

from onto_platform.ingestion.extractors.types import (
    ExtractedDocument,
    ExtractionFailure,
    Section,
)

_SAMPLE_ROW_LIMIT = 5


async def extract_xlsx(path: str, filename: str) -> ExtractedDocument:
    def _do() -> ExtractedDocument:
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            raise ExtractionFailure(f"openpyxl failed on {filename}: {e}") from e
        sections: list[Section] = []
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            headers = [
                str(c) if c is not None else "" for c in (rows[0] if rows else [])
            ]
            sample: list[tuple[Any, ...]] = rows[1 : 1 + _SAMPLE_ROW_LIMIT]
            text_lines = [", ".join(headers)]
            for r in sample:
                text_lines.append(", ".join("" if c is None else str(c) for c in r))
            sections.append(
                Section(
                    heading=ws.title,
                    text="\n".join(text_lines),
                    structured={
                        "headers": headers,
                        "sample_rows": [list(r) for r in sample],
                    },
                )
            )
        wb.close()
        return ExtractedDocument(kind="xlsx", source_filename=filename, sections=sections)

    return await asyncio.to_thread(_do)
