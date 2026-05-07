import pytest
from onto_platform.ingestion.extractors.pptx_ext import extract_pptx
from onto_platform.ingestion.extractors.types import ExtractionFailure

pytestmark = pytest.mark.asyncio


@pytest.fixture(scope="module")
def sample_pptx(tmp_path_factory):
    from pptx import Presentation
    p = tmp_path_factory.mktemp("pptx") / "sample.pptx"
    pres = Presentation()
    s1 = pres.slides.add_slide(pres.slide_layouts[1])
    s1.shapes.title.text = "Materials"
    s1.placeholders[1].text = "Materials are stocked parts."
    s2 = pres.slides.add_slide(pres.slide_layouts[1])
    s2.shapes.title.text = "Stations"
    s2.placeholders[1].text = "Stations are workshop locations."
    pres.save(str(p))
    return p


async def test_extract_pptx_per_slide(sample_pptx):
    doc = await extract_pptx(str(sample_pptx), "sample.pptx")
    assert doc.kind == "pptx"
    assert len(doc.sections) == 2
    assert doc.sections[0].heading == "Materials"
    assert "stocked parts" in doc.sections[0].text


async def test_extract_pptx_corrupt_raises(tmp_path):
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"not a pptx")
    with pytest.raises(ExtractionFailure):
        await extract_pptx(str(bad), "bad.pptx")
