import pytest
from onto_platform.ingestion.extractors.docx_ext import extract_docx
from onto_platform.ingestion.extractors.types import ExtractionFailure

pytestmark = pytest.mark.asyncio


@pytest.fixture(scope="module")
def sample_docx(tmp_path_factory):
    from docx import Document
    p = tmp_path_factory.mktemp("docx") / "sample.docx"
    d = Document()
    d.add_heading("Domain Glossary", level=1)
    d.add_paragraph("A material is a stocked part with a unique code.")
    d.add_heading("Stations", level=2)
    d.add_paragraph("Stations are workshop locations.")
    d.save(str(p))
    return p


async def test_extract_docx(sample_docx):
    doc = await extract_docx(str(sample_docx), "sample.docx")
    assert doc.kind == "docx"
    titles = [s.heading for s in doc.sections if s.heading]
    assert "Domain Glossary" in titles
    assert "Stations" in titles


async def test_extract_docx_corrupt_raises(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a docx")
    with pytest.raises(ExtractionFailure):
        await extract_docx(str(bad), "bad.docx")
